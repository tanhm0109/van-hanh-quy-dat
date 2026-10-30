"""Tính tuổi nợ, phân nhóm A–F và xuất báo cáo Excel từ sổ theo dõi vốn ứng.

    python cong-cu/tuoi_no.py <so.xlsx|so.csv> --ngay-chot 30/09/2026 --xuat <bao-cao.xlsx>
    python cong-cu/tuoi_no.py <so.xlsx> --ngay-chot 30/09/2026 --ghi-vao-so --nguoi "Nguyễn Văn A"

Mặc định KHÔNG sửa sổ gốc. --ghi-vao-so: sao lưu sổ, rồi ghi Dư nợ, Số ngày quá
hạn, Nhóm nợ vào sổ và ghi Nhật ký.
In kết quả tóm tắt dạng JSON ra màn hình để dùng khi dựng văn bản.
"""

import argparse
import json
import sys
from pathlib import Path

from so_von_ung import (
    NHOM_NO,
    NHOM_TRONG_HAN,
    TEN_SHEET_SO,
    TIEU_DE,
    LoiDuLieu,
    doc_ngay,
    doc_so,
    dinh_dang_tien,
    ghi_nhat_ky,
    kiem_tra_file_dang_mo,
    phan_nhom,
    sao_luu,
)


def tinh(dong_so, ngay_chot):
    ket_qua, canh_bao = [], []
    for d in dong_so:
        ma = d["ma_ho_so"]
        for loi in d["_loi"]:
            canh_bao.append(f"{ma} (dòng {d['_dong']}): {loi}")
        giai_ngan = d.get("da_giai_ngan") or 0
        thu_hoi = d.get("da_thu_hoi") or 0
        so_tien = d.get("so_tien_ung")
        du_no = giai_ngan - thu_hoi
        if d.get("du_no") is not None and d["du_no"] != du_no:
            canh_bao.append(
                f"{ma}: dư nợ ghi trong sổ {dinh_dang_tien(d['du_no'])} khác dư nợ tính lại "
                f"{dinh_dang_tien(du_no)} (= giải ngân − thu hồi)"
            )
        if du_no < 0:
            canh_bao.append(f"{ma}: đã thu hồi lớn hơn đã giải ngân — kiểm tra lại chứng từ")
        if so_tien is not None and giai_ngan > so_tien:
            canh_bao.append(f"{ma}: đã giải ngân lớn hơn số tiền ứng theo QĐ")
        nhom, qua_han, con_lai = phan_nhom(d.get("han_hoan_tra"), du_no, ngay_chot)
        if nhom == "?":
            canh_bao.append(f"{ma}: thiếu hạn hoàn trả — không phân nhóm được (không tự suy luận)")
        ket_qua.append(d | {"du_no_tinh": du_no, "nhom": nhom, "qua_han": qua_han, "con_lai": con_lai})
    return ket_qua, canh_bao


def tong_hop(ket_qua):
    bang = []
    for nhom, mo_ta, hanh_dong in NHOM_NO:
        ds = [r for r in ket_qua if r["nhom"] == nhom]
        bang.append({"nhom": nhom, "mo_ta": mo_ta, "hanh_dong": hanh_dong,
                     "so_khoan": len(ds), "du_no": sum(r["du_no_tinh"] for r in ds)})
    trong = [b for b in bang if b["nhom"] in NHOM_TRONG_HAN]
    qua = [b for b in bang if b["nhom"] not in NHOM_TRONG_HAN]
    return bang, {
        "trong_han": {"so_khoan": sum(b["so_khoan"] for b in trong), "du_no": sum(b["du_no"] for b in trong)},
        "qua_han": {"so_khoan": sum(b["so_khoan"] for b in qua), "du_no": sum(b["du_no"] for b in qua)},
        "thieu_han": {"so_khoan": sum(1 for r in ket_qua if r["nhom"] == "?"),
                      "du_no": sum(r["du_no_tinh"] for r in ket_qua if r["nhom"] == "?")},
        "da_tat_toan": sum(1 for r in ket_qua if r["nhom"] == "-"),
    }


def xuat_excel(dich, ket_qua, bang, tong, canh_bao, ngay_chot, nguon):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    dam = Font(bold=True)
    nen = PatternFill("solid", fgColor="D9E1F2")
    wb = Workbook()

    ws = wb.active
    ws.title = "Tong hop"
    ws.append([f"BẢNG TỔNG HỢP PHÂN LOẠI NỢ VỐN ỨNG — số liệu tính đến ngày {ngay_chot:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Nguồn: {nguon}"])
    ws.append([])
    ws.append(["Nhóm", "Tình trạng", "Số khoản", "Dư nợ (đồng)", "Hành động tối thiểu"])
    for o in ws[4]:
        o.font, o.fill = dam, nen
    for b in bang:
        ws.append([b["nhom"], b["mo_ta"], b["so_khoan"], b["du_no"], b["hanh_dong"]])
    ws.append([])
    for nhan, k in (("Cộng trong hạn (A+B)", "trong_han"), ("Cộng quá hạn (C–F)", "qua_han"),
                    ("Thiếu hạn hoàn trả — chưa phân nhóm", "thieu_han")):
        ws.append(["", nhan, tong[k]["so_khoan"], tong[k]["du_no"]])
        ws.cell(row=ws.max_row, column=2).font = dam
    ws.append(["", "Khoản đã tất toán (không tính)", tong["da_tat_toan"]])
    for r in ws.iter_rows(min_row=5, min_col=4, max_col=4):
        r[0].number_format = "#,##0"
    for chu, w in zip("ABCDE", (8, 38, 10, 20, 70)):
        ws.column_dimensions[chu].width = w

    cot = ["ma_ho_so", "don_vi_nhan_von", "du_an", "so_hop_dong", "so_tien_ung", "da_giai_ngan",
           "da_thu_hoi", "han_hoan_tra", "nguon_hoan_tra", "tinh_trang_du_an"]

    def sheet_ds(ten, ds):
        s = wb.create_sheet(ten)
        s.append([TIEU_DE[c] for c in cot] + ["Dư nợ tính lại (đồng)", "Số ngày quá hạn", "Còn lại (ngày)", "Nhóm nợ"])
        for o in s[1]:
            o.font, o.fill = dam, nen
        for r in ds:
            s.append([r.get(c) for c in cot] + [r["du_no_tinh"], r["qua_han"], r["con_lai"], r["nhom"]])
        for row in s.iter_rows(min_row=2):
            for o in row:
                if o.column in (5, 6, 7, 11):
                    o.number_format = "#,##0"
                elif o.column == 8:
                    o.number_format = "dd/mm/yyyy"
        for j in range(1, s.max_column + 1):
            s.column_dimensions[s.cell(row=1, column=j).column_letter].width = 18
        s.freeze_panes = "B2"

    sheet_ds("Chi tiet", ket_qua)
    sheet_ds("Sap den han", [r for r in ket_qua if r["nhom"] == "B"])
    sheet_ds("Qua han", sorted([r for r in ket_qua if r["nhom"] in "CDEF"], key=lambda r: -r["qua_han"]))

    cb = wb.create_sheet("Canh bao")
    cb.append(["Cảnh báo dữ liệu — cần người phụ trách kiểm tra trước khi dùng số liệu"])
    cb["A1"].font = dam
    for c in canh_bao or ["Không có cảnh báo."]:
        cb.append([c])
    cb.column_dimensions["A"].width = 120

    dich.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dich)


def ghi_vao_so(duong_dan, ket_qua, nguoi, ngay_chot):
    from openpyxl import load_workbook

    kiem_tra_file_dang_mo(duong_dan)
    ban_sao = sao_luu(duong_dan)
    wb = load_workbook(duong_dan)
    ws = wb[TEN_SHEET_SO] if TEN_SHEET_SO in wb.sheetnames else wb.worksheets[0]
    tieu_de = {str(o.value or "").strip(): o.column for o in ws[1]}
    vi_tri = {k: tieu_de.get(TIEU_DE[k]) for k in ("du_no", "so_ngay_qua_han", "nhom_no")}
    so_thay_doi = 0
    for r in ket_qua:
        moi = {"du_no": r["du_no_tinh"], "so_ngay_qua_han": r["qua_han"], "nhom_no": r["nhom"]}
        for k, v in moi.items():
            if not vi_tri[k]:
                continue
            o = ws.cell(row=r["_dong"], column=vi_tri[k])
            if o.value != v:
                ghi_nhat_ky(wb, r["ma_ho_so"], TIEU_DE[k], o.value, v,
                            f"Tính lại theo mốc {ngay_chot:%d/%m/%Y} (tuoi_no.py)", nguoi)
                o.value = v
                so_thay_doi += 1
    wb.save(duong_dan)
    return ban_sao, so_thay_doi


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("so", type=Path, help="sổ theo dõi vốn ứng (.xlsx hoặc .csv)")
    p.add_argument("--ngay-chot", required=True, help="mốc chốt số liệu, dd/mm/yyyy")
    p.add_argument("--xuat", type=Path, help="file báo cáo Excel đầu ra")
    p.add_argument("--ghi-vao-so", action="store_true", help="ghi kết quả tính vào sổ gốc (có sao lưu, nhật ký)")
    p.add_argument("--nguoi", default="", help="người thực hiện (ghi nhật ký)")
    a = p.parse_args()

    try:
        ngay_chot = doc_ngay(a.ngay_chot)
        ket_qua, canh_bao = tinh(doc_so(a.so), ngay_chot)
        bang, tong = tong_hop(ket_qua)
        if a.xuat:
            xuat_excel(a.xuat, ket_qua, bang, tong, canh_bao, ngay_chot, a.so.name)
        ghi = None
        if a.ghi_vao_so:
            if a.so.suffix.lower() != ".xlsx":
                raise LoiDuLieu("--ghi-vao-so chỉ áp dụng cho sổ .xlsx")
            ghi = ghi_vao_so(a.so, ket_qua, a.nguoi, ngay_chot)
    except LoiDuLieu as e:
        sys.exit(f"LỖI: {e}")

    tom_tat = {
        "ngay_chot": f"{ngay_chot:%d/%m/%Y}",
        "nguon": str(a.so),
        "tong_hop_theo_nhom": [{k: b[k] for k in ("nhom", "mo_ta", "so_khoan", "du_no")} for b in bang],
        **tong,
        "ds_sap_den_han": [{"ma": r["ma_ho_so"], "don_vi": r.get("don_vi_nhan_von"), "du_no": r["du_no_tinh"],
                         "han": f"{r['han_hoan_tra']:%d/%m/%Y}", "con_lai_ngay": r["con_lai"]}
                        for r in ket_qua if r["nhom"] == "B"],
        "ds_qua_han": [{"ma": r["ma_ho_so"], "don_vi": r.get("don_vi_nhan_von"), "du_no": r["du_no_tinh"],
                     "han": f"{r['han_hoan_tra']:%d/%m/%Y}", "so_ngay_qua_han": r["qua_han"], "nhom": r["nhom"]}
                    for r in sorted(ket_qua, key=lambda r: -(r["qua_han"] or 0)) if r["nhom"] in "CDEF"],
        "canh_bao": canh_bao,
    }
    if a.xuat:
        tom_tat["file_bao_cao"] = str(a.xuat)
    if ghi:
        tom_tat["sao_luu"], tom_tat["so_o_da_cap_nhat"] = str(ghi[0]), ghi[1]
    print(json.dumps(tom_tat, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
