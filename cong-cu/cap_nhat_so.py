"""Cập nhật sổ theo dõi vốn ứng: sửa trường của một khoản, hoặc thêm khoản mới.
Mỗi lần chạy đều sao lưu sổ và ghi Nhật ký kèm chứng từ.

    # Ghi nhận hoàn trả: đặt lại tổng số đã thu hồi
    python cong-cu/cap_nhat_so.py <so.xlsx> --ma UV-2026-001 --dat da_thu_hoi=5000000000 \
        --ly-do "UNC số 123 ngày 15/10/2026" --nguoi "Nguyễn Văn A"

    # Cộng thêm một lần thu hồi vào số đã có
    python cong-cu/cap_nhat_so.py <so.xlsx> --ma UV-2026-001 --cong da_thu_hoi=2000000000 --ly-do "..."

    # Thêm khoản ứng mới
    python cong-cu/cap_nhat_so.py <so.xlsx> --them --dat ma_ho_so=UV-2026-002 --dat don_vi_nhan_von="..." ... --ly-do "QĐ số ..."

Tên trường: xem COT trong so_von_ung.py. Không cho sửa cột do công cụ tính
(du_no, so_ngay_qua_han, nhom_no) — chạy tuoi_no.py --ghi-vao-so để cập nhật.
"""

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from so_von_ung import (
    COT_TINH,
    KIEU,
    TEN_SHEET_SO,
    TIEU_DE,
    LoiDuLieu,
    doc_gia_tri,
    ghi_nhat_ky,
    kiem_tra_file_dang_mo,
    sao_luu,
)


def tach(cap: str):
    if "=" not in cap:
        raise LoiDuLieu(f"sai cú pháp {cap!r}, cần dạng truong=gia_tri")
    k, v = cap.split("=", 1)
    k = k.strip()
    if k not in KIEU:
        raise LoiDuLieu(f"không có trường {k!r}")
    if k in COT_TINH:
        raise LoiDuLieu(f"{k} do công cụ tính, không sửa tay")
    return k, doc_gia_tri(k, v)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("so", type=Path)
    p.add_argument("--ma", help="mã hồ sơ cần sửa")
    p.add_argument("--them", action="store_true", help="thêm khoản mới")
    p.add_argument("--dat", action="append", default=[], help="truong=gia_tri (đặt giá trị)")
    p.add_argument("--cong", action="append", default=[], help="truong=so_tien (cộng thêm, chỉ trường số tiền)")
    p.add_argument("--ly-do", required=True, help="lý do / số hiệu chứng từ — bắt buộc")
    p.add_argument("--nguoi", default="", help="người thực hiện")
    a = p.parse_args()

    try:
        if a.them == bool(a.ma):
            raise LoiDuLieu("dùng hoặc --ma (sửa) hoặc --them (thêm mới)")
        dat = dict(tach(x) for x in a.dat)
        cong = dict(tach(x) for x in a.cong)
        for k in cong:
            if KIEU[k] != "money":
                raise LoiDuLieu(f"--cong chỉ dùng cho trường số tiền, {k} không phải")
        if not dat and not cong:
            raise LoiDuLieu("không có gì để cập nhật")

        kiem_tra_file_dang_mo(a.so)
        wb = load_workbook(a.so)
        ws = wb[TEN_SHEET_SO] if TEN_SHEET_SO in wb.sheetnames else wb.worksheets[0]
        cot = {}
        for o in ws[1]:
            for k, td in TIEU_DE.items():
                if str(o.value or "").strip() == td:
                    cot[k] = o.column
        thieu = [k for k in list(dat) + list(cong) if k not in cot]
        if thieu:
            raise LoiDuLieu(f"sổ không có cột: {', '.join(thieu)}")

        ma_da_co = {}
        dong_trong = None
        for i in range(2, ws.max_row + 1):
            v = ws.cell(row=i, column=cot["ma_ho_so"]).value
            if v:
                ma_da_co.setdefault(str(v).strip(), []).append(i)
            elif dong_trong is None:
                dong_trong = i

        if a.them:
            ma = dat.get("ma_ho_so")
            if not ma:
                raise LoiDuLieu("thêm mới cần --dat ma_ho_so=...")
            if ma in ma_da_co:
                raise LoiDuLieu(f"mã {ma} đã có ở dòng {ma_da_co[ma][0]}")
            dong = dong_trong or ws.max_row + 1
        else:
            ma = a.ma
            if ma not in ma_da_co:
                raise LoiDuLieu(f"không thấy mã {ma} trong sổ")
            if len(ma_da_co[ma]) > 1:
                raise LoiDuLieu(f"mã {ma} bị trùng ở các dòng {ma_da_co[ma]} — sửa sổ trước")
            dong = ma_da_co[ma][0]

        ban_sao = sao_luu(a.so)
        thay_doi = []
        for k, v in cong.items():
            o = ws.cell(row=dong, column=cot[k])
            cu = doc_gia_tri(k, o.value) or 0
            dat[k] = cu + v
        if "ngay_cap_nhat" in cot and "ngay_cap_nhat" not in dat:
            dat["ngay_cap_nhat"] = dt.date.today()
        for k, v in dat.items():
            o = ws.cell(row=dong, column=cot[k])
            cu = o.value
            if isinstance(cu, dt.datetime):
                cu = cu.date()
            if cu == v:
                continue
            o.value = v
            if KIEU[k] == "money":
                o.number_format = "#,##0"
            elif KIEU[k] == "date":
                o.number_format = "dd/mm/yyyy"
            if k != "ngay_cap_nhat":
                ghi_nhat_ky(wb, ma, TIEU_DE[k], cu, v, a.ly_do, a.nguoi)
                thay_doi.append({"truong": TIEU_DE[k], "cu": "" if cu is None else str(cu), "moi": str(v)})
        wb.save(a.so)
    except LoiDuLieu as e:
        sys.exit(f"LỖI: {e}")

    print(json.dumps({"ma_ho_so": ma, "dong": dong, "thay_doi": thay_doi, "sao_luu": str(ban_sao),
                      "luu_y": "Chạy lại tuoi_no.py để cập nhật dư nợ và nhóm nợ."}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
