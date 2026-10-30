"""Định nghĩa sổ theo dõi vốn ứng và các hàm dùng chung.

Mọi phép tính số tiền, số ngày, nhóm nợ đặt ở đây để các script dùng chung một
quy tắc. Khung phân nhóm A–F lấy theo
01-nghiep-vu-quy-phat-trien-dat/03-theo-doi-thu-hoi-von-ung/README.md —
khi Quỹ ban hành quy chế chính thức, chỉ cần sửa NHOM_NO.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

TEN_SHEET_SO = "So von ung"
TEN_SHEET_NHAT_KY = "Nhat ky"

# (khóa, tiêu đề cột trong Excel, kiểu: text | date | money | int)
COT = [
    ("ma_ho_so", "Mã hồ sơ", "text"),
    ("don_vi_nhan_von", "Đơn vị nhận vốn", "text"),
    ("du_an", "Dự án", "text"),
    ("so_qd_ung_von", "Số QĐ ứng vốn", "text"),
    ("ngay_qd", "Ngày QĐ", "date"),
    ("so_hop_dong", "Số hợp đồng", "text"),
    ("ngay_hop_dong", "Ngày hợp đồng", "date"),
    ("so_tien_ung", "Số tiền ứng (đồng)", "money"),
    ("da_giai_ngan", "Đã giải ngân (đồng)", "money"),
    ("da_thu_hoi", "Đã thu hồi (đồng)", "money"),
    ("du_no", "Dư nợ (đồng)", "money"),
    ("han_hoan_tra", "Hạn hoàn trả", "date"),
    ("so_ngay_qua_han", "Số ngày quá hạn", "int"),
    ("nhom_no", "Nhóm nợ", "text"),
    ("nguon_hoan_tra", "Nguồn hoàn trả", "text"),
    ("tinh_trang_du_an", "Tình trạng dự án", "text"),
    ("ngay_cap_nhat", "Ngày cập nhật", "date"),
    ("ghi_chu", "Ghi chú", "text"),
]
KHOA = [c[0] for c in COT]
KIEU = {c[0]: c[2] for c in COT}
TIEU_DE = {c[0]: c[1] for c in COT}
KHOA_THEO_TIEU_DE = {c[1].lower(): c[0] for c in COT} | {c[0]: c[0] for c in COT}

# Cột do công cụ tính — không nhập tay.
COT_TINH = {"du_no", "so_ngay_qua_han", "nhom_no"}

# (nhóm, mô tả, hành động tối thiểu)
NHOM_NO = [
    ("A", "Trong hạn, còn > 90 ngày", "Theo dõi thường xuyên"),
    ("B", "Trong hạn, còn ≤ 90 ngày", "Công văn nhắc trước hạn"),
    ("C", "Quá hạn ≤ 90 ngày", "Công văn đôn đốc lần 1; làm việc trực tiếp"),
    ("D", "Quá hạn 91 – 180 ngày", "Công văn đôn đốc lần 2; kiểm tra hiện trường; báo cáo HĐQL"),
    ("E", "Quá hạn 181 – 365 ngày", "Báo cáo UBND tỉnh; đề xuất biện pháp xử lý; tạm dừng xem xét ứng vốn mới"),
    ("F", "Quá hạn > 365 ngày / khó thu hồi", "Báo cáo UBND tỉnh và HĐQL; đề xuất phương án xử lý theo quy định"),
]
NHOM_TRONG_HAN = {"A", "B"}


class LoiDuLieu(ValueError):
    pass


def doc_ngay(v) -> dt.date | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise LoiDuLieu(f"không đọc được ngày: {v!r} (dùng dd/mm/yyyy hoặc yyyy-mm-dd)")


def doc_tien(v) -> int | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    s = re.sub(r"[\s.,đ]|dong|đồng", "", str(v).strip().lower())
    if not re.fullmatch(r"-?\d+", s):
        raise LoiDuLieu(f"không đọc được số tiền: {v!r}")
    return int(s)


def doc_gia_tri(khoa: str, v):
    kieu = KIEU[khoa]
    if kieu == "date":
        return doc_ngay(v)
    if kieu in ("money", "int"):
        return doc_tien(v)
    return None if v is None else str(v).strip()


def phan_nhom(han: dt.date | None, du_no: int, ngay_chot: dt.date):
    """Trả về (nhóm, số ngày quá hạn, số ngày còn lại). Nhóm '-' = đã tất toán, '?' = thiếu hạn."""
    if du_no <= 0:
        return "-", 0, None
    if han is None:
        return "?", None, None
    con_lai = (han - ngay_chot).days
    if con_lai >= 0:
        return ("A" if con_lai > 90 else "B"), 0, con_lai
    qua = -con_lai
    if qua <= 90:
        nhom = "C"
    elif qua <= 180:
        nhom = "D"
    elif qua <= 365:
        nhom = "E"
    else:
        nhom = "F"
    return nhom, qua, con_lai


def doc_so(duong_dan: Path) -> list[dict]:
    """Đọc sổ (.xlsx hoặc .csv) thành danh sách dòng {khóa: giá trị, '_dong': số dòng}."""
    duong_dan = Path(duong_dan)
    if duong_dan.suffix.lower() == ".csv":
        import csv

        with duong_dan.open(encoding="utf-8-sig", newline="") as f:
            hang = list(csv.reader(f))
    else:
        from openpyxl import load_workbook

        wb = load_workbook(duong_dan, data_only=True, read_only=True)
        ws = wb[TEN_SHEET_SO] if TEN_SHEET_SO in wb.sheetnames else wb.worksheets[0]
        hang = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    if not hang:
        raise LoiDuLieu("sổ không có dữ liệu")
    tieu_de = [KHOA_THEO_TIEU_DE.get(str(h or "").strip().lower()) for h in hang[0]]
    if "ma_ho_so" not in tieu_de:
        raise LoiDuLieu("không thấy cột 'Mã hồ sơ' ở dòng tiêu đề")
    ket_qua = []
    for i, r in enumerate(hang[1:], start=2):
        dong = {"_dong": i, "_loi": []}
        for k, v in zip(tieu_de, r):
            if k is None:
                continue
            try:
                dong[k] = doc_gia_tri(k, v)
            except LoiDuLieu as e:
                dong[k] = None
                dong["_loi"].append(f"{TIEU_DE[k]}: {e}")
        if not dong.get("ma_ho_so"):
            continue
        ket_qua.append(dong)
    return ket_qua


def kiem_tra_file_dang_mo(duong_dan: Path):
    """Excel tạo file khóa '~$<tên>' khi đang mở — không ghi đè lúc người khác đang sửa."""
    duong_dan = Path(duong_dan)
    khoa = duong_dan.with_name("~$" + duong_dan.name)
    khoa_lo = duong_dan.with_name(".~lock." + duong_dan.name + "#")
    if khoa.exists() or khoa_lo.exists():
        raise LoiDuLieu(f"file đang được mở trên Excel/LibreOffice ({duong_dan.name}) — đóng file rồi chạy lại")


def sao_luu(duong_dan: Path) -> Path:
    import shutil

    duong_dan = Path(duong_dan)
    thu_muc = duong_dan.parent / "sao-luu"
    thu_muc.mkdir(exist_ok=True)
    dich = thu_muc / f"{duong_dan.stem}-{dt.datetime.now():%Y%m%d-%H%M%S}{duong_dan.suffix}"
    shutil.copy2(duong_dan, dich)
    return dich


def ghi_nhat_ky(wb, ma: str, truong: str, cu, moi, ly_do: str, nguoi: str):
    if TEN_SHEET_NHAT_KY not in wb.sheetnames:
        ws = wb.create_sheet(TEN_SHEET_NHAT_KY)
        ws.append(["Thời điểm", "Mã hồ sơ", "Trường", "Giá trị cũ", "Giá trị mới", "Lý do / chứng từ", "Người thực hiện"])
    ws = wb[TEN_SHEET_NHAT_KY]
    ws.append([dt.datetime.now().strftime("%d/%m/%Y %H:%M"), ma, truong, _chuoi(cu), _chuoi(moi), ly_do, nguoi])


def _chuoi(v):
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    return "" if v is None else v


def dinh_dang_tien(n: int | None) -> str:
    return "" if n is None else f"{n:,}".replace(",", ".")
