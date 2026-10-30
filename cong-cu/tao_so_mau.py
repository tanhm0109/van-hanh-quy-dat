"""Tạo sổ theo dõi vốn ứng mẫu (.xlsx) — trống, có định dạng, sheet Nhật ký và Hướng dẫn.

    python cong-cu/tao_so_mau.py bieu-mau/01-nghiep-vu-von-ung/so-theo-doi-von-ung.xlsx
"""

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from so_von_ung import COT, COT_TINH, NHOM_NO, TEN_SHEET_NHAT_KY, TEN_SHEET_SO

DO_RONG = {"text": 22, "date": 13, "money": 18, "int": 10}
SO_DONG_CO_SAN = 200


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("dich", type=Path)
    p.add_argument("--ghi-de", action="store_true", help="cho phép ghi đè file đã có")
    a = p.parse_args()
    if a.dich.exists() and not a.ghi_de:
        sys.exit(f"Đã có {a.dich} — thêm --ghi-de nếu muốn tạo lại (dữ liệu cũ sẽ mất).")

    wb = Workbook()
    ws = wb.active
    ws.title = TEN_SHEET_SO
    dam = Font(bold=True)
    nen_tieu_de = PatternFill("solid", fgColor="D9E1F2")
    nen_tinh = PatternFill("solid", fgColor="EDEDED")

    for j, (khoa, tieu_de, kieu) in enumerate(COT, start=1):
        o = ws.cell(row=1, column=j, value=tieu_de)
        o.font = dam
        o.fill = nen_tinh if khoa in COT_TINH else nen_tieu_de
        o.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        chu = get_column_letter(j)
        ws.column_dimensions[chu].width = DO_RONG[kieu]
        for i in range(2, SO_DONG_CO_SAN + 2):
            o = ws.cell(row=i, column=j)
            if kieu == "money":
                o.number_format = "#,##0"
            elif kieu == "date":
                o.number_format = "dd/mm/yyyy"
            if khoa in COT_TINH:
                o.fill = nen_tinh
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"

    nk = wb.create_sheet(TEN_SHEET_NHAT_KY)
    nk.append(["Thời điểm", "Mã hồ sơ", "Trường", "Giá trị cũ", "Giá trị mới", "Lý do / chứng từ", "Người thực hiện"])
    for o in nk[1]:
        o.font = dam
        o.fill = nen_tieu_de
    for chu, w in zip("ABCDEFG", (17, 16, 20, 18, 18, 40, 18)):
        nk.column_dimensions[chu].width = w

    hd = wb.create_sheet("Huong dan")
    dong = [
        "SỔ THEO DÕI VỐN ỨNG — HƯỚNG DẪN",
        "",
        "1. Mỗi khoản ứng một dòng ở sheet 'So von ung'. Mã hồ sơ không trùng nhau.",
        "2. Cột nền xám (Dư nợ, Số ngày quá hạn, Nhóm nợ) do công cụ tính — không nhập tay.",
        "   Dư nợ = Đã giải ngân − Đã thu hồi.",
        "3. Ngày nhập dạng dd/mm/yyyy. Số tiền nhập số nguyên, đơn vị đồng.",
        "4. Mọi thay đổi qua công cụ cap_nhat_so.py được ghi vào sheet 'Nhat ky' kèm chứng từ.",
        "   Sửa tay trực tiếp thì tự ghi một dòng vào 'Nhat ky'.",
        "5. Chỉ người chủ sổ được sửa; người khác mở chế độ chỉ đọc.",
        "6. Đối chiếu số dư với kế toán trước mỗi kỳ báo cáo.",
        "",
        "KHUNG PHÂN NHÓM NỢ (đề xuất — thay bằng quy chế chính thức khi có):",
    ] + [f"   {n} — {m}: {h}" for n, m, h in NHOM_NO]
    for d in dong:
        hd.append([d])
    hd["A1"].font = Font(bold=True, size=13)
    hd.column_dimensions["A"].width = 110

    a.dich.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.dich)
    print(f"Đã tạo {a.dich}")


if __name__ == "__main__":
    main()
