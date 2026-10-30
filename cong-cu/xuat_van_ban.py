"""Chuyển dự thảo Markdown (.md) sang Word (.docx), tùy chọn thêm PDF.

    python cong-cu/xuat_van_ban.py <du-thao.md>                 # → du-thao.docx cùng thư mục
    python cong-cu/xuat_van_ban.py <du-thao.md> --pdf           # thêm du-thao.pdf (cần LibreOffice)
    python cong-cu/xuat_van_ban.py <du-thao.md> --bo-muc-nhap   # bỏ mục "THÔNG TIN CÒN THIẾU…" ở cuối

Hỗ trợ phần Markdown dùng trong biểu mẫu của repo: tiêu đề #, đoạn văn, danh
sách, bảng |…|, **đậm**, *nghiêng*, khối ``` (giữ nguyên dòng), > trích dẫn.
Chỗ trống {{…}} và […] được tô vàng để người soạn dễ thấy.

Định dạng trang theo 00-tai-lieu-chung/quy-uoc-van-ban.md (A4, Times New Roman).
Thông số lề/cỡ chữ để ở THE_THUC — cần đối chiếu quy định về thể thức văn bản
hiện hành [CẦN KIỂM TRA] trước khi dùng chính thức.
"""

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

THE_THUC = {
    "font": "Times New Roman",
    "co_chu": 14,
    "co_chu_bang": 12,
    "le_tren": 2.0, "le_duoi": 2.0, "le_trai": 3.0, "le_phai": 1.5,  # cm
    "thut_dau_dong": 1.0,  # cm
    "cach_sau_doan": 6,  # pt
}

MUC_NHAP = re.compile(r"^#{1,3}\s*THÔNG TIN CÒN THIẾU", re.I)
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|\{\{.*?\}\}|\[…\]|\[\.\.\.\]|\[[^\]]+\]\([^)]+\))")


def dat_font(run, co=None, dam=None, nghieng=None, mono=False):
    ten = "Courier New" if mono else THE_THUC["font"]
    run.font.name = ten
    run._element.rPr.rFonts.set(qn("w:eastAsia"), ten)
    run.font.size = Pt(co or THE_THUC["co_chu"])
    if dam is not None:
        run.bold = dam
    if nghieng is not None:
        run.italic = nghieng


def them_chu(p, text, co=None, dam=False, nghieng=False):
    """Thêm text có định dạng inline vào đoạn p."""
    for phan in INLINE.split(text):
        if not phan:
            continue
        d, n, to_vang = dam, nghieng, False
        if phan.startswith("**") and phan.endswith("**") and len(phan) > 4:
            them_chu(p, phan[2:-2], co, True, nghieng)
            continue
        if phan.startswith("*") and phan.endswith("*") and len(phan) > 2:
            them_chu(p, phan[1:-1], co, dam, True)
            continue
        if phan.startswith("`") and phan.endswith("`"):
            phan = phan[1:-1]
        elif m := re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", phan):
            phan = m.group(1)
        if re.search(r"\{\{.*?\}\}|\[…\]|\[\.\.\.\]", phan):
            to_vang = True
        r = p.add_run(phan)
        dat_font(r, co, d, n)
        if to_vang:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def doan(doc, text="", can=None, thut=True, **kw):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(THE_THUC["cach_sau_doan"])
    pf.space_before = Pt(0)
    if thut:
        pf.first_line_indent = Cm(THE_THUC["thut_dau_dong"])
    p.alignment = can if can is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    if text:
        them_chu(p, text, **kw)
    return p


def bang(doc, dong_md):
    hang = []
    for d in dong_md:
        o = [c.strip() for c in d.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in o if c):
            continue
        hang.append(o)
    if not hang:
        return
    so_cot = max(len(h) for h in hang)
    t = doc.add_table(rows=0, cols=so_cot)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(hang):
        cells = t.add_row().cells
        for j in range(so_cot):
            c = cells[j]
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            them_chu(c.paragraphs[0], h[j] if j < len(h) else "", co=THE_THUC["co_chu_bang"], dam=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def chuyen(md: str, bo_muc_nhap: bool) -> Document:
    doc = Document()
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.top_margin, s.bottom_margin = Cm(THE_THUC["le_tren"]), Cm(THE_THUC["le_duoi"])
        s.left_margin, s.right_margin = Cm(THE_THUC["le_trai"]), Cm(THE_THUC["le_phai"])
    st = doc.styles["Normal"]
    st.font.name = THE_THUC["font"]
    st.element.rPr.rFonts.set(qn("w:eastAsia"), THE_THUC["font"])
    st.font.size = Pt(THE_THUC["co_chu"])

    dong = md.splitlines()
    bat_dau_khoi = re.compile(r"^\s*([#|>]|```|[-*+]\s|\d+[.)]\s)")

    def noi_tiep(i, dau):
        """Gộp các dòng xuống hàng (không mở khối mới) vào nội dung đang có."""
        khoi = [dau]
        while i + 1 < len(dong) and dong[i + 1].strip() and not bat_dau_khoi.match(dong[i + 1]) \
                and not re.fullmatch(r"-{3,}", dong[i + 1].strip()):
            i += 1
            khoi.append(dong[i].strip())
        return i, " ".join(khoi)

    i = 0
    while i < len(dong):
        d = dong[i]
        s = d.strip()
        if bo_muc_nhap and (MUC_NHAP.match(s) or (s == "---" and any(MUC_NHAP.match(x.strip()) for x in dong[i + 1:i + 3]))):
            break
        if s.startswith("```"):
            i += 1
            while i < len(dong) and not dong[i].strip().startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(dong[i])
                dat_font(r, co=11, mono=True)
                i += 1
            i += 1
            continue
        if s.startswith("|"):
            khoi = []
            while i < len(dong) and dong[i].strip().startswith("|"):
                khoi.append(dong[i])
                i += 1
            bang(doc, khoi)
            continue
        if not s or re.fullmatch(r"-{3,}|\*{3,}", s):
            i += 1
            continue
        if m := re.match(r"^(#{1,6})\s+(.*)", s):
            cap = len(m.group(1))
            if cap == 1:
                doan(doc, m.group(2).upper(), can=WD_ALIGN_PARAGRAPH.CENTER, thut=False, dam=True)
            elif cap == 2:
                doan(doc, m.group(2), can=WD_ALIGN_PARAGRAPH.CENTER if i < 6 else None, thut=False, dam=True)
            else:
                doan(doc, m.group(2), thut=False, dam=True, nghieng=(cap >= 4))
        elif m := re.match(r"^(\s*)[-*+]\s+(\[[ xX]\]\s+)?(.*)", d):
            hop = ("☒ " if m.group(2).strip()[1] in "xX" else "☐ ") if m.group(2) else "– "
            i, noi_dung = noi_tiep(i, m.group(3))
            p = doan(doc, thut=False)
            p.paragraph_format.left_indent = Cm(THE_THUC["thut_dau_dong"] + len(m.group(1)) * 0.25)
            them_chu(p, hop + noi_dung)
        elif m := re.match(r"^(\s*)(\d+[.)])\s+(.*)", d):
            i, noi_dung = noi_tiep(i, m.group(3))
            p = doan(doc, thut=False)
            p.paragraph_format.left_indent = Cm(THE_THUC["thut_dau_dong"] + len(m.group(1)) * 0.25)
            them_chu(p, f"{m.group(2)} {noi_dung}")
        elif s.startswith(">"):
            doan(doc, s.lstrip("> "), nghieng=True)
        else:
            i, noi_dung = noi_tiep(i, s)
            doan(doc, noi_dung)
        i += 1
    return doc


def sang_pdf(docx: Path) -> Path:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        for c in (r"C:\Program Files\LibreOffice\program\soffice.exe", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
            if Path(c).exists():
                exe = c
    if not exe:
        raise RuntimeError("không tìm thấy LibreOffice để xuất PDF — cài LibreOffice, hoặc mở file Word rồi Lưu thành PDF")
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                   check=True, capture_output=True, timeout=180)
    return docx.with_suffix(".pdf")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("nguon", type=Path)
    p.add_argument("--xuat", type=Path, help="đường dẫn .docx đầu ra (mặc định cùng tên)")
    p.add_argument("--pdf", action="store_true")
    p.add_argument("--bo-muc-nhap", action="store_true", help="bỏ mục THÔNG TIN CÒN THIẾU… ở cuối dự thảo")
    a = p.parse_args()

    md = a.nguon.read_text(encoding="utf-8")
    dich = a.xuat or a.nguon.with_suffix(".docx")
    chuyen(md, a.bo_muc_nhap).save(dich)
    print(f"Word: {dich}")
    if a.pdf:
        try:
            print(f"PDF:  {sang_pdf(dich)}")
        except (RuntimeError, subprocess.SubprocessError) as e:
            print(f"Không xuất được PDF: {e}", file=sys.stderr)

    trong = re.findall(r"\{\{.*?\}\}|\[…\]|\[\.\.\.\]", md)
    if trong:
        print(f"LƯU Ý: còn {len(trong)} chỗ trống ({{{{…}}}} / […]) — đã tô vàng trong file Word.")
    if not a.bo_muc_nhap and any(MUC_NHAP.match(x.strip()) for x in md.splitlines()):
        print("LƯU Ý: văn bản còn mục 'THÔNG TIN CÒN THIẾU' — xoá trước khi trình ký, hoặc chạy lại với --bo-muc-nhap.")


if __name__ == "__main__":
    main()
